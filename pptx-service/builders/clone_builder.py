"""
CloneBuilder — builds a deck by copying template slides and replacing their
content in place, so every slide keeps the template's exact design.

Plan schema (one entry per output slide, in order):
{
  "slides": [{
    "template_slide": 7,                       # 1-based slide of the template to copy
    "texts":  {"<shape id>": "Text with **bold** spans\\nSecond paragraph"
               | ["para 1", {"text": "sub point", "level": 1}]},
    "charts": {"<shape id>": {"categories": [...], "series": [{"name": "...", "values": [...]}],
                              "number_format": "0%", "title": "..."}
                             | {"series": [{"name": "...", "points": [[x, y], ...]}]}},   # XY charts
    "tables": {"<shape id>": {"rows": [["Header", ...], ["Row", ...]]}},
    "remove": ["<shape id>", ...],
    "elements": [{"replace": "<shape id>" | "box": [x, y, w, h], "type": "waterfall", ...}]
  }],
}

Text that would not fit the space the template gives it is shrunk (down to 75%
of its size); anything still too long is reported in warnings.
"""
from __future__ import annotations

import copy
import io
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

from builders.deck_assembler import DeckAssembler
from builders.template_catalog import iter_shapes
from builders.text_fit import (NS, EMU_PER_IN, StyleResolver, body_metrics, fit_lines, lines_needed,
                               text_width_pt, strip_markup)

A_NS = NS["a"]
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_BOLD = re.compile(r"\*\*(.+?)\*\*")
MIN_SCALE = 0.75
XY_TYPES = {XL_CHART_TYPE.XY_SCATTER, XL_CHART_TYPE.XY_SCATTER_LINES, XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
            XL_CHART_TYPE.XY_SCATTER_SMOOTH, XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS}


def _a(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


def _segments(text: str) -> list[tuple[str, bool]]:
    """Split '**bold** rest' markup into (text, bold) runs."""
    out, pos = [], 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out or [("", False)]


def _norm_paragraphs(value) -> list[dict]:
    if isinstance(value, (int, float)):
        value = str(value)
    if isinstance(value, str):
        return [{"text": t, "level": None} for t in value.split("\n")]
    out = []
    for item in value or []:
        if isinstance(item, dict):
            out.append({"text": str(item.get("text", "")), "level": item.get("level")})
        else:
            out.append({"text": str(item), "level": None})
    return out or [{"text": "", "level": None}]


class CloneBuilder:
    def __init__(self, template_bytes: bytes, palette: dict | None = None, font: str | None = None):
        self.template = template_bytes
        self.palette = palette or {}
        self.font = font
        self.warnings: list[dict] = []

    # ── Public ─────────────────────────────────────────────────────────────────
    def build(self, plan: dict) -> bytes:
        slides_plan = [s for s in plan.get("slides", []) if isinstance(s, dict)]
        sequence = [int(s.get("template_slide", 0)) for s in slides_plan]
        assembled = DeckAssembler(self.template).assemble(sequence)
        prs = Presentation(io.BytesIO(assembled))
        self.styles = StyleResolver(prs)
        self.drawer = self._make_drawer(prs)
        for i, (slide, sp) in enumerate(zip(prs.slides, slides_plan), 1):
            self._fill_slide(i, slide, sp)
        out = io.BytesIO()
        prs.save(out)
        return out.getvalue()

    def _warn(self, slide: int, shape_id, kind: str, msg: str, **extra) -> None:
        self.warnings.append({"slide": slide, "shape": str(shape_id) if shape_id is not None else None,
                              "kind": kind, "message": msg, **extra})

    # ── Slide ──────────────────────────────────────────────────────────────────
    def _shape_map(self, slide) -> dict:
        return {str(sh.shape_id): (sh, box) for sh, box, _ in iter_shapes(slide.shapes)}

    def _fill_slide(self, n: int, slide, sp: dict) -> None:
        shapes = self._shape_map(slide)

        for sid, value in (sp.get("texts") or {}).items():
            hit = shapes.get(str(sid))
            if not hit:
                self._warn(n, sid, "missing", "Shape id not found on the template slide")
                continue
            sh, box = hit
            if not sh.has_text_frame:
                self._warn(n, sid, "not_text", "Shape has no text frame")
                continue
            try:
                self._set_shape_text(n, slide, sh, box, value)
            except Exception as e:
                self._warn(n, sid, "error", f"Text fill failed: {e}")

        for sid, data in (sp.get("charts") or {}).items():
            hit = shapes.get(str(sid))
            if not hit or not getattr(hit[0], "has_chart", False) or not hit[0].has_chart:
                self._warn(n, sid, "missing", "Chart not found")
                continue
            try:
                self._fill_chart(hit[0].chart, data)
            except Exception as e:
                self._warn(n, sid, "error", f"Chart fill failed: {e}")

        for sid, data in (sp.get("tables") or {}).items():
            hit = shapes.get(str(sid))
            if not hit or not getattr(hit[0], "has_table", False) or not hit[0].has_table:
                self._warn(n, sid, "missing", "Table not found")
                continue
            try:
                self._fill_table(n, hit[0], data)
            except Exception as e:
                self._warn(n, sid, "error", f"Table fill failed: {e}")

        to_remove = {str(x) for x in (sp.get("remove") or [])}
        for el in sp.get("elements") or []:
            if not isinstance(el, dict):
                continue
            box = el.get("box")
            rid = el.get("replace")
            if rid is not None:
                hit = shapes.get(str(rid))
                if not hit:
                    self._warn(n, rid, "missing", "Shape to replace not found")
                    continue
                box = hit[1]
                to_remove.add(str(rid))
            if not box or len(box) != 4:
                continue
            x, y, w, h = [float(v) for v in box]
            try:
                self.drawer._element(slide, {**el, "x": x, "y": y, "w": w, "h": h})
            except Exception as e:
                self._warn(n, rid, "error", f"Element {el.get('type')} failed: {e}")

        for sid in to_remove:
            hit = shapes.get(sid)
            if hit:
                el = hit[0].element
                if el.getparent() is not None:
                    el.getparent().remove(el)

    # ── Text ───────────────────────────────────────────────────────────────────
    def _set_shape_text(self, n, slide, sh, box, value) -> None:
        tx = sh.element.find(".//p:txBody", NS)
        bm = body_metrics(tx)
        width = max(0.05, box[2] - bm["l"] - bm["r"])
        height = max(0.05, box[3] - bm["t"] - bm["b"])
        scale = bm["font_scale"]

        old = [(self._p_text(p), self.styles.paragraph_props(sh, slide, p)) for p in tx.findall("a:p", NS)]
        old_nonempty = [(t, pr) for t, pr in old if t.strip()]
        first_props = (old_nonempty or old or [("", {"size": 18})])[0][1]
        allowed = max(fit_lines(height, first_props, scale),
                      lines_needed(old_nonempty, width, scale) if old_nonempty else 1)

        self._replace_paragraphs(tx, _norm_paragraphs(value))

        new = [(self._p_text(p), self.styles.paragraph_props(sh, slide, p)) for p in tx.findall("a:p", NS)]
        new = [(t, pr) for t, pr in new if t.strip()] or new
        if not bm["wrap"]:
            # Single-line boxes: only the width matters.
            widest = max((text_width_pt(strip_markup(t), pr["size"] * scale, pr.get("face"), pr.get("bold"), pr.get("caps"))
                          for t, pr in new), default=0)
            limit = max(width * 72, max((text_width_pt(t, pr["size"] * scale, pr.get("face"), pr.get("bold"), pr.get("caps"))
                                         for t, pr in old_nonempty), default=0))
            if widest > limit:
                s = max(MIN_SCALE, limit / widest)
                self._scale_fonts(tx, new, scale * s, bm)
                if limit / widest < MIN_SCALE:
                    self._warn(n, sh.shape_id, "overflow", "Text too wide for its box", over_by=round(widest / limit, 2))
            return

        need = lines_needed(new, width, scale)
        if need <= allowed:
            return
        s = 0.95
        while s >= MIN_SCALE - 1e-6:
            if lines_needed(new, width, scale * s) <= max(allowed, fit_lines(height, first_props, scale * s)):
                self._scale_fonts(tx, new, scale * s, bm)
                return
            s -= 0.05
        self._scale_fonts(tx, new, scale * MIN_SCALE, bm)
        text_len = sum(len(strip_markup(t)) for t, _ in new)
        self._warn(n, sh.shape_id, "overflow", "Text does not fit its box",
                   lines_needed=lines_needed(new, width, scale * MIN_SCALE), lines_allowed=allowed,
                   chars=text_len)

    @staticmethod
    def _p_text(p) -> str:
        return "".join(t.text or "" for t in p.iter(_a("t")))

    def _scale_fonts(self, tx, paras, effective_scale: float, bm: dict) -> None:
        """Set explicit sizes on every run (base size × scale), dropping autofit font scaling."""
        bp = tx.find("a:bodyPr", NS)
        na = bp.find("a:normAutofit", NS) if bp is not None else None
        if na is not None and na.get("fontScale"):
            del na.attrib["fontScale"]
            na.attrib.pop("lnSpcReduction", None)
        for p_el, (_, props) in zip([p for p in tx.findall("a:p", NS) if self._p_text(p).strip()] or tx.findall("a:p", NS), paras):
            size = max(6.0, round(props["size"] * effective_scale * 2) / 2)
            for r in p_el.findall("a:r", NS) + p_el.findall("a:fld", NS):
                rpr = r.find("a:rPr", NS)
                if rpr is None:
                    rpr = etree.SubElement(r, _a("rPr"))
                    r.insert(0, rpr)
                rpr.set("sz", str(int(size * 100)))
            end = p_el.find("a:endParaRPr", NS)
            if end is not None:
                end.set("sz", str(int(size * 100)))

    def _replace_paragraphs(self, tx, paras: list[dict]) -> None:
        olds = tx.findall("a:p", NS)
        protos = [p for p in olds if self._p_text(p).strip()] or olds

        def level_of(p):
            ppr = p.find("a:pPr", NS)
            return int(ppr.get("lvl", 0)) if ppr is not None else 0

        new_ps = []
        for i, para in enumerate(paras):
            if para["level"] is not None:
                same = [p for p in protos if level_of(p) == int(para["level"])]
                proto = same[min(i, len(same) - 1)] if same else protos[min(i, len(protos) - 1)]
            else:
                proto = protos[min(i, len(protos) - 1)]
            new_ps.append(self._paragraph_from(proto, para))
        for p in olds:
            tx.remove(p)
        for p in new_ps:
            tx.append(p)

    def _paragraph_from(self, proto, para: dict):
        p = copy.deepcopy(proto)
        runs = p.findall("a:r", NS)
        base = runs[0].find("a:rPr", NS) if runs else None
        if base is None:
            end = p.find("a:endParaRPr", NS)
            base = etree.Element(_a("rPr"), **dict(end.attrib)) if end is not None else etree.Element(_a("rPr"))
            if end is not None:
                for child in end:
                    base.append(copy.deepcopy(child))
        bold_run = next((r.find("a:rPr", NS) for r in runs if r.find("a:rPr", NS) is not None
                         and r.find("a:rPr", NS).get("b") == "1"), None)
        for child in list(p):
            if child.tag in (_a("r"), _a("br"), _a("fld")):
                p.remove(child)
        if para["level"] is not None:
            ppr = p.find("a:pPr", NS)
            if ppr is None:
                ppr = etree.Element(_a("pPr"))
                p.insert(0, ppr)
            if int(para["level"]) > 0:
                ppr.set("lvl", str(int(para["level"])))
        end = p.find("a:endParaRPr", NS)
        base_bold = base.get("b") == "1"
        for text, bold in _segments(para["text"]):
            if not text:
                continue
            r = etree.Element(_a("r"))
            if bold and not base_bold:
                rpr = copy.deepcopy(bold_run) if bold_run is not None else copy.deepcopy(base)
                rpr.set("b", "1")
            else:
                rpr = copy.deepcopy(base)
            rpr.attrib.pop("dirty", None)
            r.append(rpr)
            t = etree.SubElement(r, _a("t"))
            t.text = text
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        return p

    # ── Charts ─────────────────────────────────────────────────────────────────
    def _fill_chart(self, chart, data: dict) -> None:
        series = [s for s in data.get("series", []) if isinstance(s, dict)]
        if not series:
            raise ValueError("no series")
        fmt = data.get("number_format")
        # Without an explicit format keep the template's (e.g. "0.0%" on a donut).
        src_fmt = chart._chartSpace.find(f".//{{{C_NS}}}val//{{{C_NS}}}formatCode")
        cache_fmt = fmt or (src_fmt.text if src_fmt is not None and src_fmt.text else None)
        if chart.chart_type in XY_TYPES:
            cd = XyChartData()
            for s in series:
                ser = cd.add_series(str(s.get("name", "")))
                for pt in s.get("points", []):
                    if isinstance(pt, (list, tuple)) and len(pt) >= 2:
                        ser.add_data_point(float(pt[0]), float(pt[1]))
        else:
            cats = [str(c) for c in data.get("categories", [])]
            cd = CategoryChartData(number_format=cache_fmt) if cache_fmt else CategoryChartData()
            cd.categories = cats
            for s in series:
                vals = [None if v is None else float(v) for v in s.get("values", [])][:len(cats)]
                vals += [None] * (len(cats) - len(vals))
                cd.add_series(str(s.get("name", "")), vals)
        chart.replace_data(cd)
        self._apply_point_colors(chart, series)
        self._fit_value_axis(chart, series)
        if fmt:
            for plot in chart.plots:
                if plot.has_data_labels:
                    plot.data_labels.number_format = fmt
                    plot.data_labels.number_format_is_linked = False
        if data.get("title") and chart.has_title and chart.chart_title.has_text_frame:
            tf = chart.chart_title.text_frame
            if tf.paragraphs and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].text = str(data["title"])
                for r in tf.paragraphs[0].runs[1:]:
                    r.text = ""

    @staticmethod
    def _apply_point_colors(chart, series: list[dict]) -> None:
        """series[i]["point_colors"] = {"<point index>": "RRGGBB"} recolors single points
        (e.g. the company's own bar); {} clears all per-point colors. Absent keeps the template's."""
        sers = chart._chartSpace.findall(f".//{{{C_NS}}}ser")
        for ser_el, spec in zip(sers, series):
            wanted = spec.get("point_colors")
            if wanted is None or not isinstance(wanted, dict):
                continue
            dpts = ser_el.findall(f"{{{C_NS}}}dPt")
            proto = next((d for d in dpts if d.find(f"{{{C_NS}}}spPr") is not None), None)
            for d in dpts:
                ser_el.remove(d)
            anchor = None
            for tag in ("marker", "invertIfNegative", "spPr", "tx", "order"):
                anchor = ser_el.find(f"{{{C_NS}}}{tag}")
                if anchor is not None:
                    break
            base_sppr = ser_el.find(f"{{{C_NS}}}spPr")
            for idx in sorted((int(k) for k in wanted), reverse=True):
                color = str(wanted.get(str(idx), wanted.get(idx, ""))).lstrip("#")
                if len(color) != 6:
                    continue
                if proto is not None:
                    d = copy.deepcopy(proto)
                else:
                    d = etree.Element(f"{{{C_NS}}}dPt")
                    etree.SubElement(d, f"{{{C_NS}}}idx")
                    etree.SubElement(d, f"{{{C_NS}}}invertIfNegative", val="0")
                    etree.SubElement(d, f"{{{C_NS}}}bubble3D", val="0")
                    d.append(copy.deepcopy(base_sppr) if base_sppr is not None else etree.Element(f"{{{C_NS}}}spPr"))
                d.find(f"{{{C_NS}}}idx").set("val", str(idx))
                sp = d.find(f"{{{C_NS}}}spPr")
                for old in sp.findall("a:solidFill", NS) + sp.findall("a:pattFill", NS) + sp.findall("a:gradFill", NS) + sp.findall("a:noFill", NS):
                    sp.remove(old)
                sf = etree.Element(_a("solidFill"))
                etree.SubElement(sf, _a("srgbClr"), val=color.upper())
                sp.insert(0, sf)
                if anchor is not None:
                    anchor.addnext(d)
                else:
                    ser_el.insert(0, d)

    @staticmethod
    def _fit_value_axis(chart, series: list[dict]) -> None:
        """Release fixed axis bounds that would clip or flatten the new data."""
        try:
            axis = chart.value_axis
        except Exception:
            return  # pie / doughnut
        vals = [float(v) for s in series for v in s.get("values", []) if v is not None]
        if not vals:
            return
        stacked = "STACKED" in str(chart.chart_type)
        top = max(vals)
        if stacked:
            n = max(len(s.get("values", [])) for s in series)
            top = max(sum(float(s["values"][i]) for s in series if i < len(s.get("values", [])) and s["values"][i] is not None and float(s["values"][i]) > 0) for i in range(n))
        mx, mn = axis.maximum_scale, axis.minimum_scale
        if mx is not None and (top > mx or top < mx * 0.45):
            axis.maximum_scale = None
        if mn is not None and min(vals) < mn:
            axis.minimum_scale = None

    # ── Tables ─────────────────────────────────────────────────────────────────
    def _fill_table(self, n: int, frame, data: dict) -> None:
        rows = [[("" if c is None else str(c)) for c in r] for r in data.get("rows", []) if isinstance(r, list)]
        if not rows:
            raise ValueError("no rows")
        tbl = frame.table
        tbl_el = tbl._tbl
        merged = any(c.is_merge_origin or c.is_spanned for r in tbl.rows for c in r.cells)
        n_cols = max(len(r) for r in rows)
        grid = tbl_el.find("a:tblGrid", NS)
        cols = grid.findall("a:gridCol", NS)
        total_w = sum(int(c.get("w")) for c in cols)

        if n_cols != len(cols) and not merged:
            if n_cols > len(cols):
                for _ in range(n_cols - len(cols)):
                    grid.append(copy.deepcopy(cols[-1]))
                    for tr in tbl_el.findall("a:tr", NS):
                        tr.append(copy.deepcopy(tr.findall("a:tc", NS)[-1]))
            else:
                for gc in cols[n_cols:]:
                    grid.remove(gc)
                for tr in tbl_el.findall("a:tr", NS):
                    for tc in tr.findall("a:tc", NS)[n_cols:]:
                        tr.remove(tc)
            cols = grid.findall("a:gridCol", NS)
            widths = [int(c.get("w")) for c in cols]
            factor = total_w / max(1, sum(widths))
            for c, w in zip(cols, widths):
                c.set("w", str(int(w * factor)))
        elif n_cols != len(cols):
            self._warn(n, frame.shape_id, "table_shape", f"Table has merged cells; kept {len(cols)} columns")

        trs = tbl_el.findall("a:tr", NS)
        orig_h = sum(int(tr.get("h")) for tr in trs)
        if len(rows) > len(trs):
            proto = trs[-2] if len(trs) >= 3 else trs[-1]
            for _ in range(len(rows) - len(trs)):
                new_tr = copy.deepcopy(proto)
                tbl_el.findall("a:tr", NS)[-1].addnext(new_tr)
        elif len(rows) < len(trs):
            for tr in trs[len(rows):]:
                tbl_el.remove(tr)
        trs = tbl_el.findall("a:tr", NS)

        # Keep the table inside the template's footprint when rows were added.
        if len(trs) > 1 and sum(int(tr.get("h")) for tr in trs) > orig_h * 1.05:
            head = int(trs[0].get("h"))
            body_h = max(int(0.2 * EMU_PER_IN), int((orig_h * 1.05 - head) / (len(trs) - 1)))
            for tr in trs[1:]:
                tr.set("h", str(body_h))

        for r_i, tr in enumerate(trs):
            tcs = tr.findall("a:tc", NS)
            for c_i, tc in enumerate(tcs):
                if tc.get("hMerge") == "1" or tc.get("vMerge") == "1":
                    continue
                value = rows[r_i][c_i] if c_i < len(rows[r_i]) else ""
                tx = tc.find("a:txBody", NS)
                if tx is None:
                    continue
                self._replace_paragraphs(tx, _norm_paragraphs(value))
        frame.height = Emu(sum(int(tr.get("h")) for tr in trs))

    # ── Engine-drawn elements (charts the template doesn't have) ──────────────
    def _make_drawer(self, prs):
        from builders.pptx_builder import PptxBuilder, PALETTE, DEFAULT_FONT
        d = PptxBuilder.__new__(PptxBuilder)
        d.prs = prs
        d.PALETTE = {**PALETTE, **{k: v for k, v in self.palette.items() if v}}
        d.DEFAULT_FONT = self.font or DEFAULT_FONT
        d.warnings = []
        return d
