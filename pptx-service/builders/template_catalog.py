"""
TemplateCatalog — describes every slide of a template deck so a model can reuse
the slides as layouts: each shape with its id, kind, absolute position, current
text, effective font and how much text it can hold; charts with their type and
data shape; tables with their cells; pictures with a small thumbnail.
"""
from __future__ import annotations

import base64
import hashlib
import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from builders.template_profiler import TemplateProfiler
from builders.text_fit import (NS, EMU_PER_IN, StyleResolver, body_metrics, chars_per_line, fit_lines,
                               lines_needed)

THUMB_PX = 160
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _r(v: float) -> float:
    return round(v, 2)


def iter_shapes(shapes, transform=None, parent_id=None):
    """Yields (shape, (x, y, w, h) in inches on the slide, parent group id), recursing into groups."""
    t = transform or (lambda x, y, w, h: (x, y, w, h))
    for sh in shapes:
        try:
            x, y, w, h = (sh.left or 0), (sh.top or 0), (sh.width or 0), (sh.height or 0)
        except Exception:
            continue
        ax, ay, aw, ah = t(x, y, w, h)
        yield sh, (ax / EMU_PER_IN, ay / EMU_PER_IN, aw / EMU_PER_IN, ah / EMU_PER_IN), parent_id
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = sh.element.find("p:grpSpPr/a:xfrm", NS)
            if xfrm is None:
                continue
            off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
            choff, chext = xfrm.find("a:chOff", NS), xfrm.find("a:chExt", NS)
            if None in (off, ext, choff, chext):
                continue
            ox, oy = int(off.get("x")), int(off.get("y"))
            cx, cy = int(ext.get("cx")), int(ext.get("cy"))
            chox, choy = int(choff.get("x")), int(choff.get("y"))
            chcx, chcy = int(chext.get("cx")) or 1, int(chext.get("cy")) or 1
            sx, sy = cx / chcx, cy / chcy

            def child_t(x, y, w, h, _t=t, ox=ox, oy=oy, chox=chox, choy=choy, sx=sx, sy=sy):
                return _t(ox + (x - chox) * sx, oy + (y - choy) * sy, w * sx, h * sy)

            yield from iter_shapes(sh.shapes, child_t, sh.shape_id)


def fill_color(sp_pr, theme_colors: dict) -> str | None:
    """Hex of a solid fill (srgbClr, or schemeClr resolved through the theme)."""
    if sp_pr is None:
        return None
    sf = sp_pr.find("a:solidFill", NS)
    if sf is None:
        return None
    el = sf.find("a:srgbClr", NS)
    base = el.get("val") if el is not None else None
    if el is None:
        el = sf.find("a:schemeClr", NS)
        if el is None:
            return None
        key = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}.get(el.get("val"), el.get("val"))
        base = theme_colors.get(key)
    return apply_color_mods(base, el) if base else None


def apply_color_mods(hex_color: str, el) -> str:
    """Apply DrawingML lumMod / lumOff / tint / shade children to a hex color."""
    import colorsys
    r, g, b = (int(hex_color[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for m in el:
        name, val = m.tag.split("}")[-1], int(m.get("val", 100000)) / 100000
        if name == "lumMod":
            l *= val
        elif name == "lumOff":
            l += val
        elif name == "tint":
            l = l * val + (1 - val)
        elif name == "shade":
            l *= val
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "".join(f"{round(c * 255):02X}" for c in (r, g, b))


def shape_kind(sh) -> str:
    tag = sh.element.tag.split("}")[-1]
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    if tag == "pic":
        return "picture"
    if tag == "cxnSp":
        return "line"
    if tag == "graphicFrame":
        return "other"
    if sh.has_text_frame and (sh.text_frame.text.strip() or sh.is_placeholder):
        return "text"
    return "shape"


class TemplateCatalog:
    def __init__(self, template_bytes: bytes):
        self.raw = template_bytes
        self.prs = Presentation(io.BytesIO(template_bytes))
        self.styles = StyleResolver(self.prs)
        self.media: dict[str, str] = {}

    def extract(self, with_thumbnails: bool = True) -> dict:
        profile = TemplateProfiler(self.raw).extract()
        self.theme_colors = profile.get("colors", {})
        slides = [self._slide(i, s, with_thumbnails) for i, s in enumerate(self.prs.slides, 1)]
        faces: dict[str, int] = {}
        for s in slides:
            for sh in s["shapes"]:
                f = (sh.get("font") or {}).get("face")
                if f:
                    faces[f] = faces.get(f, 0) + len(sh.get("text", ""))
        return {
            "slide_width_in": _r(self.prs.slide_width / EMU_PER_IN),
            "slide_height_in": _r(self.prs.slide_height / EMU_PER_IN),
            "theme_fonts": self.styles.theme_fonts,
            "body_font": max(faces, key=faces.get) if faces else self.styles.theme_fonts.get("minor"),
            "colors": profile.get("colors", {}),
            "palette": profile.get("palette", {}),
            "slides": slides,
            "media": self.media if with_thumbnails else {},
        }

    # ── Per slide ──────────────────────────────────────────────────────────────
    def _slide(self, index: int, slide, with_thumbnails: bool) -> dict:
        shapes = []
        for sh, box, parent in iter_shapes(slide.shapes):
            try:
                shapes.append(self._shape(slide, sh, box, parent, with_thumbnails))
            except Exception as e:  # never let one odd shape break the catalog
                shapes.append({"id": sh.shape_id, "kind": "other", "name": sh.name, "error": str(e)[:120],
                               "box": [_r(v) for v in box]})
        return {
            "index": index,
            "layout": slide.slide_layout.name,
            "hidden": slide.element.get("show") == "0",
            "shapes": shapes,
        }

    def _shape(self, slide, sh, box, parent, with_thumbnails) -> dict:
        kind = shape_kind(sh)
        d: dict = {"id": sh.shape_id, "kind": kind, "name": sh.name, "box": [_r(v) for v in box]}
        if parent is not None:
            d["group"] = parent
        if sh.is_placeholder:
            ph = sh.element.find(".//p:nvPr/p:ph", NS)
            d["placeholder"] = ph.get("type", "body") if ph is not None else "body"
        if kind in ("text", "shape") and sh.has_text_frame:
            self._text(slide, sh, box, d)
        elif kind == "chart":
            d["chart"] = self._chart(sh.chart)
        elif kind == "table":
            d["table"] = self._table(slide, sh)
        elif kind == "picture":
            d["picture"] = self._picture(sh, with_thumbnails)
        return d

    def _text(self, slide, sh, box, d: dict) -> None:
        tf = sh.text_frame
        tx = sh.element.find(".//p:txBody", NS)
        bm = body_metrics(tx)
        paras = []
        for p_el in tx.findall("a:p", NS):
            text = "".join(t.text or "" for t in p_el.iter(f"{{{NS['a']}}}t"))
            props = self.styles.paragraph_props(sh, slide, p_el)
            paras.append((text, props))
        if any(p_el.find("a:fld", NS) is not None for p_el in tx.findall("a:p", NS)):
            fld = tx.find(".//a:fld", NS)
            d["field"] = fld.get("type") if fld is not None else "field"
        text = tf.text
        if not text.strip() and not sh.is_placeholder:
            return
        first = next((pr for t, pr in paras if t.strip()), paras[0][1] if paras else {"size": 18})
        scale = bm["font_scale"]
        width = max(0.05, box[2] - bm["l"] - bm["r"])
        height = max(0.05, box[3] - bm["t"] - bm["b"])
        d["text"] = text[:400]
        d["font"] = {"size": _r(first["size"] * scale), "bold": first.get("bold", False),
                     "caps": first.get("caps", False), "face": first.get("face")}
        if len(paras) > 1:
            d["paragraphs"] = [{"text": t[:120], "level": pr["level"], "size": _r(pr["size"] * scale),
                                "bold": pr.get("bold", False)} for t, pr in paras[:12] if t.strip()]
        if bm["vertical"]:
            d["vertical"] = True
        # Designers often draw boxes shorter than their text and let it spill,
        # so the line budget is what fits OR what the template already uses.
        used = lines_needed([(t, pr) for t, pr in paras if t.strip()], width, scale) if text.strip() else 0
        d["max_lines"] = max(fit_lines(height, first, scale), used)
        d["chars_per_line"] = chars_per_line(width, first, scale)

    def _chart(self, chart) -> dict:
        out: dict = {}
        try:
            out["type"] = str(chart.chart_type).split(".")[-1].split(" ")[0]
        except Exception:
            out["type"] = "unknown"
        try:
            plots = list(chart.plots)
            out["plots"] = len(plots)
            plot = plots[0]
            try:
                out["categories"] = [str(c) for c in plot.categories][:24]
            except Exception:
                pass
            out["series"] = []
            for pl in plots:
                for s in pl.series:
                    try:
                        vals = [v for v in s.values][:24]
                    except Exception:
                        vals = []
                    entry = {"name": s.name, "values": [round(v, 4) if isinstance(v, float) else v for v in vals]}
                    base = fill_color(s._element.find(f"{{{C_NS}}}spPr"), self.theme_colors)
                    if base:
                        entry["color"] = base
                    # Points the designer colored individually (e.g. the company's own bar)
                    points = {}
                    for d in s._element.findall(f"{{{C_NS}}}dPt"):
                        col = fill_color(d.find(f"{{{C_NS}}}spPr"), self.theme_colors)
                        if col:
                            points[d.find(f"{{{C_NS}}}idx").get("val")] = col
                    if points and "DOUGHNUT" not in out["type"] and "PIE" not in out["type"]:
                        entry["point_colors"] = points
                    out["series"].append(entry)
            if plot.has_data_labels:
                try:
                    out["label_format"] = plot.data_labels.number_format
                except Exception:
                    pass
        except Exception:
            pass
        try:
            out["axis_format"] = chart.value_axis.tick_labels.number_format
        except Exception:
            pass
        out["legend"] = bool(chart.has_legend)
        out["title"] = chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else None
        return out

    def _table(self, slide, sh) -> dict:
        tbl = sh.table
        rows = []
        merged = False
        for r in tbl.rows:
            row = []
            for c in r.cells:
                if c.is_merge_origin or c.is_spanned:
                    merged = True
                row.append(c.text[:60])
            rows.append(row)
        size = None
        try:
            cell = tbl.cell(min(1, len(tbl.rows) - 1), 0)
            p_el = cell._tc.find(".//a:p", NS)
            if p_el is not None:
                r = p_el.find("a:r/a:rPr", NS)
                size = int(r.get("sz")) / 100 if r is not None and r.get("sz") else None
        except Exception:
            pass
        return {"rows": len(tbl.rows), "cols": len(tbl.columns), "cells": rows[:16], "merged": merged,
                "font_size": size}

    def _picture(self, sh, with_thumbnails) -> dict:
        out: dict = {}
        try:
            img = sh.image
            key = hashlib.sha1(img.blob).hexdigest()[:12]
            out["media"] = key
            out["format"] = img.ext
            descr = sh.element.find(".//p:cNvPr", NS)
            if descr is not None and descr.get("descr"):
                out["alt"] = descr.get("descr")[:120]
            if with_thumbnails and key not in self.media and img.ext.lower() in ("png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"):
                from PIL import Image
                im = Image.open(io.BytesIO(img.blob))
                im.thumbnail((THUMB_PX, THUMB_PX))
                if im.mode not in ("RGB", "L"):
                    bg = Image.new("RGB", im.size, (255, 255, 255))
                    bg.paste(im.convert("RGBA"), mask=im.convert("RGBA").split()[-1])
                    im = bg
                buf = io.BytesIO()
                im.convert("RGB").save(buf, "JPEG", quality=70)
                self.media[key] = base64.b64encode(buf.getvalue()).decode()
        except Exception:
            pass
        return out
